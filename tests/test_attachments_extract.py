"""``swarm_builder.attachments.extract``: bounded, safe reads of uploaded files.

The assertions here are deliberately about *content*, not about "it returned
something": a feature whose whole value is that the model sees the right column
names and thresholds has to be tested against an exact expected table, or the
test would pass with an empty extraction.

Fixtures are built at test time (``openpyxl`` and ``python-pptx`` write them)
rather than committed as binaries: a committed workbook is opaque in review, and
these tests need to state precisely which cells they assert on.
"""

from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path

import pytest

from swarm_builder.attachments import extract as extract_module
from swarm_builder.attachments.limits import (
    MAX_ATTACHMENT_CHARS,
    MAX_CELL_CHARS,
    MAX_COLS,
    MAX_INFLATED_BYTES,
    MAX_ROWS_PER_SHEET,
    MAX_SHEETS,
    MAX_SLIDES,
)
from swarm_builder.attachments.models import (
    AttachmentReadError,
    UnsupportedAttachmentError,
)


def _extract(name: str, data: bytes, **kwargs: object) -> extract_module.Extraction:
    return extract_module.extract(name, data, **kwargs)  # type: ignore[arg-type]


def _workbook_bytes(sheets: dict[str, list[list[object]]]) -> bytes:
    """Build an .xlsx in memory from {sheet name: rows}."""
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _deck_bytes(slides: list[dict[str, object]]) -> bytes:
    """Build a .pptx in memory; each slide takes title/body/notes/table."""
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    presentation = pptx.Presentation()
    layout = presentation.slide_layouts[5]  # title only
    for spec in slides:
        slide = presentation.slides.add_slide(layout)
        if spec.get("title"):
            slide.shapes.title.text = str(spec["title"])
        if spec.get("body"):
            box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
            box.text_frame.text = str(spec["body"])
        table_spec = spec.get("table")
        if isinstance(table_spec, list):
            rows, cols = len(table_spec), len(table_spec[0])
            shape = slide.shapes.add_table(rows, cols, Inches(1), Inches(3), Inches(6), Inches(1))
            for r, row in enumerate(table_spec):
                for c, value in enumerate(row):
                    shape.table.cell(r, c).text = str(value)
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(spec["notes"])
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Detection and refusals
# ---------------------------------------------------------------------------


def test_detection_is_by_extension_and_ignores_a_lying_content_type() -> None:
    kind, media_type = extract_module.detect_kind("data.csv", "application/x-evil")
    assert kind == "csv"
    assert media_type == "text/csv"
    # A text file renamed .xlsx is read as a workbook (and fails), never sniffed
    # into the CSV reader because the browser claimed text/csv.
    kind, _ = extract_module.detect_kind("books.xlsx", "text/csv")
    assert kind == "xlsx"


def test_an_unknown_suffix_names_the_fix() -> None:
    with pytest.raises(UnsupportedAttachmentError) as raised:
        extract_module.detect_kind("legacy.xls", "application/vnd.ms-excel")
    message = str(raised.value)
    assert "re-save it as .xlsx" in message
    assert ".csv" in message


def test_a_known_media_type_is_used_only_when_the_suffix_is_unknown() -> None:
    kind, media_type = extract_module.detect_kind("download", "image/png")
    assert (kind, media_type) == ("image", "image/png")


def test_an_empty_file_is_refused() -> None:
    with pytest.raises(AttachmentReadError) as raised:
        _extract("empty.csv", b"")
    assert "empty" in str(raised.value)


def test_a_file_of_nul_bytes_is_refused() -> None:
    with pytest.raises(AttachmentReadError) as raised:
        _extract("binary.csv", b"\x00" * 64)
    assert "no readable text" in str(raised.value)


# ---------------------------------------------------------------------------
# CSV / TSV
# ---------------------------------------------------------------------------


def test_csv_becomes_an_exact_markdown_table() -> None:
    extraction = _extract("orders.csv", b"Order ID,Region,Amount\n1,EMEA,10\n2,APAC,20\n")
    assert extraction.truncated is False
    assert extraction.notes == []
    assert extraction.text == (
        "| Order ID | Region | Amount |\n"
        "| --- | --- | --- |\n"
        "| 1 | EMEA | 10 |\n"
        "| 2 | APAC | 20 |"
    )


def test_csv_escapes_pipes_and_collapses_embedded_newlines() -> None:
    data = b'id,note\n1,"a | b"\n2,"line one\nline two"\n'
    extraction = _extract("notes.csv", data)
    assert r"| 1 | a \| b |" in extraction.text
    assert "| 2 | line one line two |" in extraction.text
    assert "\n\n" not in extraction.text.split("\n", 2)[2]


def test_csv_cell_longer_than_the_cap_is_ellipsized() -> None:
    long_value = "x" * (MAX_CELL_CHARS + 50)
    extraction = _extract("wide.csv", f"k\n{long_value}\n".encode())
    cell = extraction.text.splitlines()[-1]
    assert cell.endswith("… |")
    assert len(cell) < MAX_CELL_CHARS + 20


def test_semicolon_delimited_and_bom_prefixed_csv_still_parses() -> None:
    extraction = _extract("euro.csv", "\ufeffa;b\n1;2\n".encode("utf-8"))
    assert "| a | b |" in extraction.text
    assert "| 1 | 2 |" in extraction.text
    assert extraction.notes == []


def test_cp1252_csv_is_decoded_and_the_note_says_so() -> None:
    extraction = _extract("legacy.csv", "café,prix\n1,2\n".encode("cp1252"))
    assert "café" in extraction.text
    assert any("decoded as cp1252" in note for note in extraction.notes)


def test_csv_row_cap_is_reported() -> None:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["n"])
    for value in range(MAX_ROWS_PER_SHEET + 25):
        writer.writerow([value])
    extraction = _extract("big.csv", buffer.getvalue().encode())
    assert extraction.truncated is True
    assert f"only the first {MAX_ROWS_PER_SHEET} rows" in " ".join(extraction.notes)
    # header row + separator + the capped rows, all read (the header counts
    # toward the cap, which is what the note says too)
    assert len(extraction.text.splitlines()) == MAX_ROWS_PER_SHEET + 1


def test_csv_column_cap_is_reported_and_applied() -> None:
    header = ",".join(f"c{index}" for index in range(MAX_COLS + 5))
    row = ",".join("1" for _ in range(MAX_COLS + 5))
    extraction = _extract("wide.csv", f"{header}\n{row}\n".encode())
    assert extraction.truncated is True
    assert f"c{MAX_COLS - 1}" in extraction.text
    assert f"c{MAX_COLS}" not in extraction.text


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------


def test_xlsx_two_sheets_are_rendered_exactly() -> None:
    data = _workbook_bytes(
        {
            "Orders": [["Order ID", "Region"], [1, "EMEA"], [2, "APAC"]],
            "Rules": [["Region", "Limit"], ["EMEA", 100]],
        }
    )
    extraction = _extract("pipeline.xlsx", data)
    assert extraction.truncated is False
    assert extraction.text == (
        '### Sheet "Orders" (3 rows x 2 cols)\n'
        "| Order ID | Region |\n"
        "| --- | --- |\n"
        "| 1 | EMEA |\n"
        "| 2 | APAC |\n"
        "\n"
        '### Sheet "Rules" (2 rows x 2 cols)\n'
        "| Region | Limit |\n"
        "| --- | --- |\n"
        "| EMEA | 100 |"
    )


def test_xlsx_sheet_cap_is_reported() -> None:
    data = _workbook_bytes({f"S{index}": [["a"], [index]] for index in range(MAX_SHEETS + 3)})
    extraction = _extract("many.xlsx", data)
    assert extraction.truncated is True
    assert f"only the first {MAX_SHEETS} of {MAX_SHEETS + 3} sheets were read" in extraction.notes
    assert "S0" in extraction.text
    assert f"S{MAX_SHEETS}" not in extraction.text


def test_xlsx_row_cap_is_reported() -> None:
    rows = [["n"]] + [[index] for index in range(MAX_ROWS_PER_SHEET + 40)]
    extraction = _extract("long.xlsx", _workbook_bytes({"S": rows}))
    assert extraction.truncated is True
    assert any(f"{MAX_ROWS_PER_SHEET} rows" in note for note in extraction.notes)


def test_xlsx_formulas_are_included_when_their_values_are_not_cached() -> None:
    """A workbook whose limits live in formulas is the case this feature is for.

    ``openpyxl`` saves no cached results, so the value pass reads those cells as
    empty; the bounded second pass must hand the model the formula text instead.
    """
    data = _workbook_bytes(
        {"Thresholds": [["Region", "Limit"], ["EMEA", "=100*1.2"], ["APAC", "=80+5"]]}
    )
    extraction = _extract("thresholds.xlsx", data)
    assert any("computes values with formulas" in note for note in extraction.notes)
    assert "### Formulas" in extraction.text
    assert "Thresholds!B2 =100*1.2" in extraction.text
    assert "Thresholds!B3 =80+5" in extraction.text


def test_xlsx_without_formulas_has_no_formula_block() -> None:
    extraction = _extract("plain.xlsx", _workbook_bytes({"S": [["a"], [1]]}))
    assert "### Formulas" not in extraction.text


def test_a_text_file_renamed_xlsx_is_an_unreadable_workbook() -> None:
    with pytest.raises(AttachmentReadError) as raised:
        _extract("fake.xlsx", b"this is not a zip container")
    message = str(raised.value)
    assert "not a readable" in message
    assert "fake.xlsx" in message
    assert "/" not in message  # never a path


def test_an_oversized_inflation_is_refused_before_any_parser_runs() -> None:
    """A small upload that declares a huge uncompressed size must not be read."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        chunk = b"\x00" * (1024 * 1024)
        with archive.open("xl/worksheets/sheet1.xml", "w") as member:
            for _ in range((MAX_INFLATED_BYTES // len(chunk)) + 2):
                member.write(chunk)
    payload = buffer.getvalue()
    assert len(payload) < 1024 * 1024  # compressed to almost nothing
    with pytest.raises(AttachmentReadError) as raised:
        _extract("bomb.xlsx", payload)
    assert "expands to more than" in str(raised.value)


def test_xlsx_long_text_output_is_clipped_and_reported() -> None:
    rows = [["text"]] + [
        [f"row {index} " + "y" * MAX_CELL_CHARS] for index in range(MAX_ROWS_PER_SHEET)
    ]
    extraction = _extract("huge.xlsx", _workbook_bytes({"S": rows}))
    assert len(extraction.text) <= MAX_ATTACHMENT_CHARS
    assert extraction.truncated is True
    assert any("characters were kept" in note for note in extraction.notes)


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------


def test_pptx_renders_titles_bodies_tables_and_notes() -> None:
    data = _deck_bytes(
        [
            {
                "title": "Intake",
                "body": "Tickets arrive from email",
                "notes": "Ask about the SLA here",
                "table": [["Field", "Type"], ["id", "string"]],
            },
            {"title": "Routing"},
        ]
    )
    extraction = _extract("process.pptx", data)
    assert "## Slide 1 — Intake" in extraction.text
    assert "Tickets arrive from email" in extraction.text
    assert "| Field | Type |" in extraction.text
    assert "Notes: Ask about the SLA here" in extraction.text
    assert "## Slide 2 — Routing" in extraction.text
    assert extraction.truncated is False


def test_pptx_slide_cap_is_reported() -> None:
    data = _deck_bytes([{"title": f"S{index}"} for index in range(MAX_SLIDES + 2)])
    extraction = _extract("long.pptx", data)
    assert extraction.truncated is True
    assert f"only the first {MAX_SLIDES} of {MAX_SLIDES + 2} slides were read" in extraction.notes


def test_a_text_file_renamed_pptx_is_an_unreadable_deck() -> None:
    with pytest.raises(AttachmentReadError) as raised:
        _extract("fake.pptx", b"not a presentation at all")
    assert "not a readable" in str(raised.value)


# ---------------------------------------------------------------------------
# Images
# ---------------------------------------------------------------------------


def test_an_image_yields_no_text_and_says_where_it_goes() -> None:
    extraction = _extract("diagram.png", _real_image("PNG"), declared_media_type="image/png")
    assert extraction.text == ""
    assert extraction.truncated is False
    assert extraction.notes == [extract_module.IMAGE_NOTE]


def test_supported_summary_lists_every_extension() -> None:
    summary = extract_module.supported_summary()
    for suffix in (".csv", ".tsv", ".xlsx", ".pptx", ".png", ".jpeg", ".webp", ".gif"):
        assert suffix in summary


def test_extract_is_deterministic_for_the_same_input() -> None:
    data = _workbook_bytes({"S": [["a", "b"], [1, 2]]})
    first = _extract("same.xlsx", data)
    second = _extract("same.xlsx", data)
    assert first == second


def test_filenames_in_errors_are_sanitized() -> None:
    """A traversal-ish name must not appear verbatim in a user-facing message."""
    with pytest.raises(UnsupportedAttachmentError) as raised:
        extract_module.detect_kind("../../etc/passwd.xls", None)
    message = str(raised.value)
    assert "/" not in message
    assert "\\" not in message
    assert "\x00" not in message


def test_relative_and_absolute_paths_do_not_change_detection() -> None:
    assert extract_module.detect_kind("dir/data.csv", None)[0] == "csv"
    assert Path("dir/data.csv").suffix == ".csv"


# ---------------------------------------------------------------------------
# Image validation
# ---------------------------------------------------------------------------


def _real_image(fmt: str = "PNG") -> bytes:
    """A genuinely decodable image, via Pillow (which ships with python-pptx)."""
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (8, 6), "white").save(buffer, format=fmt)
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("name", "fmt", "media_type"),
    [
        ("shot.png", "PNG", "image/png"),
        ("shot.jpg", "JPEG", "image/jpeg"),
        ("shot.gif", "GIF", "image/gif"),
    ],
)
def test_a_real_image_is_accepted(name: str, fmt: str, media_type: str) -> None:
    extraction = _extract(name, _real_image(fmt), declared_media_type=media_type)
    assert extraction.text == ""
    assert extraction.notes == [extract_module.IMAGE_NOTE]


def test_a_text_file_renamed_png_is_refused_at_extraction() -> None:
    """The common user mistake: catching it here beats a provider 400 later."""
    with pytest.raises(AttachmentReadError) as raised:
        _extract("shot.png", b"this is not an image at all", declared_media_type="image/png")
    assert "does not look like a PNG image" in str(raised.value)


def test_a_truncated_png_is_refused_at_extraction() -> None:
    """A valid header is not enough: the provider would reject the rest.

    Pillow reports this instance as "not a readable image" (it cannot identify a
    header-only file) rather than "truncated"; both are app-authored sentences,
    and what matters is that the file never reaches the provider.
    """
    with pytest.raises(AttachmentReadError) as raised:
        _extract("shot.png", _real_image("PNG")[:40], declared_media_type="image/png")
    message = str(raised.value)
    assert "shot.png" in message
    assert "image" in message
    assert "/" not in message


def _png_with_declared_size(width: int, height: int) -> bytes:
    """A valid PNG whose IHDR claims ``width`` x ``height`` (CRC recomputed).

    The CRC has to be right, or Pillow refuses the file as unidentifiable before
    its decompression-bomb guard ever runs -- which is a different test.
    """
    import struct
    import zlib

    payload = bytearray(_real_image("PNG"))
    payload[16:20] = struct.pack(">I", width)
    payload[20:24] = struct.pack(">I", height)
    payload[29:33] = struct.pack(">I", zlib.crc32(bytes(payload[12:29])) & 0xFFFFFFFF)
    return bytes(payload)


def test_an_image_declaring_an_absurd_pixel_count_is_refused() -> None:
    """Pillow's decompression-bomb guard is part of the defence, not a detail."""
    with pytest.raises(AttachmentReadError) as raised:
        _extract(
            "bomb.png",
            _png_with_declared_size(100_000, 100_000),
            declared_media_type="image/png",
        )
    assert "too many pixels" in str(raised.value)
